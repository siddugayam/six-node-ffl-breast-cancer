#!/usr/bin/env python3
"""
analyses/six_node_followups, Q9: the filter counts per DISTINCT (miRNA, TF) pair instead of per workbook row.
Read from the author zip (Hypergeometric_Test_2.xlsx; value.txt).  A pair is the (miRNA label, TF label) of a row;
the one TF that Excel converted (DEC1: a date in the left block, the serial 37226 in the right block, "Dec-01" in
value.txt) is keyed by its Excel serial so that its rows match across the three.  Retained = "Yes" in the right-hand
block.  Shared targets i, Nmir and Ntf are those of the left block (= value.txt); expected = Nmir*Ntf/N with N = 256;
upper tail = P(X >= i) (scipy hypergeom.sf(i - 1, 256, Nmir, Ntf)), defined only when Ntf <= 256 and Nmir <= 256.
usage: q9_distinct_pairs.py <author zip>
"""
import sys, os, io, zipfile, datetime, collections
import openpyxl
from scipy.stats import hypergeom
H = os.path.dirname(os.path.abspath(__file__)); N = 256
z = zipfile.ZipFile(sys.argv[1]); ws = openpyxl.load_workbook(io.BytesIO(z.read('Hypergeometric_Test_2.xlsx')), data_only=True).worksheets[0]
def key(tf):
    if isinstance(tf, datetime.datetime): return str((tf - datetime.datetime(1899, 12, 30)).days)   # Excel serial
    return str(tf)
L = [(str(ws.cell(r, 2).value), key(ws.cell(r, 3).value), ws.cell(r, 4).value, ws.cell(r, 5).value, ws.cell(r, 6).value) for r in range(3, 2679)]
R = [(str(ws.cell(r, 11).value), key(ws.cell(r, 12).value), ws.cell(r, 15).value == 'Yes') for r in range(3, 2679)]
out = []
P = lambda *a: (out.append(' '.join(str(x) for x in a)), print(*a))
vals = collections.defaultdict(set)
for m, t, a, b, c in L: vals[(m, t)].add((a, b, c))
P(f'Q9 rows {len(L)}; distinct tested pairs {len(vals)}; pairs whose repeated rows carry different (Nmir, Ntf, i): {sum(1 for v in vals.values() if len(v) > 1)}')
ret = {(m, t) for m, t, y in R if y}; notret = {(m, t) for m, t, y in R if not y}
P(f'Q9 distinct retained pairs {len(ret)} (from {sum(1 for x in R if x[2])} retained rows); distinct pairs with rows of both flags {len(ret & notret)}; '
  f'distinct retained pairs matched to value.txt / the left block {len(ret & set(vals))}')
M = [(k, next(iter(vals[k]))) for k in sorted(ret & set(vals))]
i0 = sum(1 for k, (a, b, c) in M if c == 0); i1 = sum(1 for k, (a, b, c) in M if c >= 1); gt = sum(1 for k, (a, b, c) in M if c > a * b / N)
P(f'Q9 distinct retained pairs: share no target (i = 0) {i0}; share >= 1 {i1}; share more than expected (Nmir*Ntf/N) {gt}')
valid = {k: v for k, vv in vals.items() for v in [next(iter(vv))] if v[1] <= N and v[0] <= N}
up = {k for k, (a, b, c) in valid.items() if hypergeom.sf(c - 1, N, a, b) < 0.05}
P(f'Q9 distinct tested pairs with upper-tail P < 0.05 at N = 256: {len(up)} of {len(valid)} with a defined P ({len(vals) - len(valid)} pairs have Ntf > 256)')
big = {k for k, vv in vals.items() if next(iter(vv))[1] > N}
P(f'Q9 Ntf > 256 ("#NUM!"): rows {sum(1 for m, t, a, b, c in L if b > N)}; distinct pairs {len(big)}; distinct TF entries {len({t for m, t in big})}: '
  f'{dict(collections.Counter(next(iter(vals[k]))[1] for k in big))} (Ntf value: pairs)')
ntf = {l.split('\t')[0]: int(l.split('\t')[1]) for l in z.read('Ntf.txt').decode().split('\n') if l.strip()}
byt = collections.defaultdict(set)
for k in big: byt[k[1]].add(next(iter(vals[k]))[1])
P('Q9 the TF entries with Ntf > 256 (Ntf used; Ntf.txt value or "absent"): ' + '; '.join(f'{t} {sorted(v)} ({ntf.get(t, "absent")})' for t, v in sorted(byt.items())))
open(f'{H}/q9_distinct_pairs.txt', 'w').write('\n'.join(out) + '\n')
