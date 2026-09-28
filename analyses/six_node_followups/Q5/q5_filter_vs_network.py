#!/usr/bin/env python3
"""
analyses/six_node_followups, Q5: does the miRNA-TF pair filter determine the network, and where does TF<->miRNA reciprocity
come from?  Read-only on the project.

Label mapping (the project's own harmonisation, scripts/01_network_assembly/02_build_canonical_network.py):
  1. data/name_map.json merge_map (the 803 raw SIF labels -> 587 canonical nodes), when the workbook label is one of them;
  2. otherwise canon2() of that script, verbatim: hsa-mir-X -> hsa-miR-X, and a -1/-2 precursor copy suffix of
     hsa-miR-<n><letter> is dropped when the unsuffixed canonical name is a network node;
  3. otherwise data/mirna_id_map.tsv (mature IDs -> canonical).
  A label is mapped only if the result is a network node of the right type.  The same rule is applied to TF labels
  (identity or merge_map).  Sensitivity (reported separately, NOT the project's rule): let-7 precursor copies
  (hsa-let-7a-1, -2, -3) collapsed to hsa-let-7a when that node exists.
Workbook: the author zip (Hypergeometric_Test_2.xlsx), right-hand BH block: retained = "Yes"; tested = every row.
Network: data/canonical_edges.tsv without the 30 legacy miRNA_miRNA edges (6,829 arcs), data/canonical_nodes.tsv.
Evidence per direction: data/layer_TF_miRNA.tsv (TransmiR, PMIDs) for TF->miRNA; data/edge_evidence_tier.tsv
(databases, support types, tier) for miRNA->TF.  TransmiR human: data/db/transmir_hsa.tsv.
usage: q5_filter_vs_network.py <author zip>
"""
import sys, os, io, re, csv, json, zipfile, collections
import openpyxl
H = os.path.dirname(os.path.abspath(__file__)); REV = '/path/to/revision'
out = []
P = lambda *a: (out.append(' '.join(str(x) for x in a)), print(*a))
# ---------------------------------------------------------------- network
nodes = {r['name']: r['type'] for r in csv.DictReader(open(f'{REV}/data/canonical_nodes.tsv'), delimiter='\t')}
E = [r for r in csv.DictReader(open(f'{REV}/data/canonical_edges.tsv'), delimiter='\t') if r['edge_type'] != 'miRNA_miRNA']
TM = {(r['source'], r['target']) for r in E if r['edge_type'] == 'TF_miRNA'}                               # TF -> miRNA
MT = {(r['source'], r['target']) for r in E if r['edge_type'] == 'miRNA_target' and nodes[r['target']] == 'TF'}  # miRNA -> TF
REC = {(t, m) for (t, m) in TM if (m, t) in MT}
mir_nodes = sorted(n for n, t in nodes.items() if t == 'miRNA'); tf_nodes = sorted(n for n, t in nodes.items() if t == 'TF')
P(f'network: {len(E)} arcs (legacy miRNA_miRNA excluded); miRNA nodes {len(mir_nodes)}, TF-typed nodes {len(tf_nodes)}; '
  f'TF->miRNA arcs {len(TM)}, miRNA->TF arcs {len(MT)}, reciprocal TF<->miRNA pairs {len(REC)}')
# ---------------------------------------------------------------- the project's harmonisation
merge_map = json.load(open(f'{REV}/data/name_map.json'))['merge_map']
idmap = {r['raw_id']: r['canonical_id'] for r in csv.DictReader(open(f'{REV}/data/mirna_id_map.tsv'), delimiter='\t')}
is_mir = lambda n: bool(re.match(r'^hsa-(mir|miR|let)-', n, flags=re.I))
def canon(n):                                           # verbatim from scripts/01_network_assembly/02_build_canonical_network.py
    if not is_mir(n): return n
    m = re.match(r'^hsa-(mir|miR)-(.+)$', n)
    if m: return 'hsa-miR-' + m.group(2)
    return n.replace('hsa-LET-', 'hsa-let-')
canon_names = set(nodes)
def canon2(n):                                          # verbatim, with canon_names = the network's canonical nodes
    c = canon(n); m = re.match(r'^(hsa-miR-\d+[a-z]?)-([12])$', c)
    if m and m.group(1) in canon_names: return m.group(1)
    return c
def mapm(label, sens=False):
    for c in (merge_map.get(label), canon2(label), idmap.get(label)):
        if c and nodes.get(c) == 'miRNA': return c
    if sens:
        m = re.match(r'^(hsa-let-7[a-z]?)-\d$', label)
        if m and nodes.get(m.group(1)) == 'miRNA': return m.group(1)
    return None
def mapt(label):
    for c in (merge_map.get(label), label):
        if c and nodes.get(c) == 'TF': return c
    return None
# ---------------------------------------------------------------- workbook
z = zipfile.ZipFile(sys.argv[1]); ws = openpyxl.load_workbook(io.BytesIO(z.read('Hypergeometric_Test_2.xlsx')), data_only=True).worksheets[0]
rows = [(str(ws.cell(r, 11).value), str(ws.cell(r, 12).value), ws.cell(r, 15).value == 'Yes') for r in range(3, 2679)]
wm = sorted({m for m, t, y in rows}); wt = sorted({t for m, t, y in rows})
for sens in (False, True):
    tag = 'sensitivity (let-7 copies collapsed; not the project rule)' if sens else "project rule"
    Mm = {m: mapm(m, sens) for m in wm}; Mt = {t: mapt(t) for t in wt}
    st = collections.defaultdict(lambda: 'none')
    for m, t, y in rows:
        cm, ct = Mm[m], Mt[t]
        if cm is None or ct is None: continue
        k = (cm, ct); st[k] = 'retained' if (y or st[k] == 'retained') else 'tested, not retained'
    nst = {}
    for (cm, ct), s in st.items():
        for n in (cm, ct):
            if s == 'retained' or nst.get(n) != 'retained': nst[n] = s if nst.get(n) != 'retained' else 'retained'
    if not sens:
        P(f'workbook labels ({tag}): miRNA labels {len(wm)}, mapped to a network miRNA node {sum(1 for v in Mm.values() if v)}; '
          f'TF labels {len(wt)}, mapped to a network TF node {sum(1 for v in Mt.values() if v)}; unmapped miRNA labels in retained rows: '
          f'{sorted({m for m, t, y in rows if y and not Mm[m]})}')
        with open(f'{H}/q5_label_map.csv', 'w', newline='') as fh:
            w = csv.writer(fh); w.writerow(['workbook_label', 'kind', 'network_node', 'in_retained_row'])
            for m in wm: w.writerow([m, 'miRNA', Mm[m] or '', any(y for mm, t, y in rows if mm == m)])
            for t in wt: w.writerow([t, 'TF', Mt[t] or '', any(y for m, tt, y in rows if tt == t)])
    for kind, NL in (('miRNA nodes', mir_nodes), ('TF-typed nodes', tf_nodes)):
        c = collections.Counter(nst.get(n, 'in no tested pair') for n in NL)
        P(f'Q5(a) {kind} ({len(NL)}), {tag}: in a retained pair {c["retained"]}; in a tested but not retained pair {c["tested, not retained"]}; '
          f'in no tested pair {c["in no tested pair"]}')
    for kind, S in (('TF->miRNA arcs', {(m, t) for (t, m) in TM}), ('miRNA->TF arcs', MT), ('reciprocal pairs', {(m, t) for (t, m) in REC})):
        c = collections.Counter(st.get(k, 'none') for k in S)
        P(f'Q5(b) {kind} ({len(S)}), {tag}: retained pair {c["retained"]}; tested, not retained {c["tested, not retained"]}; no tested pair {c["none"]}')
    if not sens:
        with open(f'{H}/q5_node_status.csv', 'w', newline='') as fh:
            w = csv.writer(fh); w.writerow(['node', 'type', 'filter_status'])
            for n in mir_nodes + tf_nodes: w.writerow([n, nodes[n], nst.get(n, 'in no tested pair')])
        with open(f'{H}/q5_arc_status.csv', 'w', newline='') as fh:
            w = csv.writer(fh); w.writerow(['source', 'target', 'direction', 'reciprocal', 'filter_status'])
            for (t, m) in sorted(TM): w.writerow([t, m, 'TF->miRNA', (t, m) in REC, st.get((m, t), 'none')])
            for (m, t) in sorted(MT): w.writerow([m, t, 'miRNA->TF', (t, m) in REC, st.get((m, t), 'none')])
# ---------------------------------------------------------------- (c) evidence per direction
lay = collections.defaultdict(list)
for r in csv.DictReader(open(f'{REV}/data/layer_TF_miRNA.tsv'), delimiter='\t'): lay[(r['source'], r['target'])].append(r)
tier = {(r['source'], r['target']): r for r in csv.DictReader(open(f'{REV}/data/edge_evidence_tier.tsv'), delimiter='\t')}
tm_rec = sum(1 for k in TM if k in lay and any(x['pmid'] not in ('', 'NA') for x in lay[k]))
P(f'Q5(c) TF->miRNA arcs with a TransmiR record (layer_TF_miRNA.tsv, PMID): {tm_rec} of {len(TM)}; of the reciprocal pairs: '
  f'{sum(1 for k in REC if k in lay)} of {len(REC)}')
ct = collections.Counter(tier[k]['tier'] if k in tier else 'no record' for k in MT)
cr = collections.Counter(tier[(m, t)]['tier'] if (m, t) in tier else 'no record' for (t, m) in REC)
db = collections.Counter(tier[(m, t)]['databases'] if (m, t) in tier else 'no record' for (t, m) in REC)
P(f'Q5(c) miRNA->TF arcs by miRNA-target evidence (edge_evidence_tier.tsv): {dict(ct)}; of the reciprocal pairs: {dict(cr)}; '
  f'databases of the reciprocal miRNA->TF arcs: {dict(db.most_common(8))}')
both = sum(1 for (t, m) in REC if (t, m) in lay and (m, t) in tier and tier[(m, t)]['tier'] != 'predicted_only')
P(f'Q5(c) reciprocal pairs with a TransmiR record for TF->miRNA AND an experimental miRNA-target record (tier strong or weak) '
  f'for miRNA->TF: {both} of {len(REC)}')
# selection: TransmiR pairs among network nodes, by whether the network has the reverse miRNA->TF arc
tr = [l.rstrip('\n').split('\t') for l in open(f'{REV}/data/db/transmir_hsa.tsv') if l.strip()]
TR = set()
for r in tr:
    cm, ct_ = mapm(r[1]), mapt(r[0])
    if cm and ct_: TR.add((ct_, cm))
with_rev = {k for k in TR if (k[1], k[0]) in MT}; no_rev = TR - with_rev
P(f'Q5(c) TransmiR TF->miRNA pairs between network TF and miRNA nodes (project rule): {len(TR)}; with the reverse miRNA->TF arc in the '
  f'network {len(with_rev)}, of which kept as TF->miRNA arcs {len(with_rev & TM)}; without it {len(no_rev)}, of which kept {len(no_rev & TM)}; '
  f'network TF->miRNA arcs not in TransmiR (project rule) {len(TM - TR)}')
npairs = len(tf_nodes) * len(mir_nodes); other = npairs - len(TR); rev_other = len(MT) - len({(m, t) for (t, m) in with_rev} & MT)
P(f'Q5(c) base rate: miRNA->TF arcs on TransmiR pairs {len(with_rev)} of {len(TR)} ({100 * len(with_rev) / len(TR):.1f} %) against {rev_other} of the other '
  f'{other} TF-miRNA node pairs ({100 * rev_other / other:.2f} %)')
# per deposited SIF file (the six Cytoscape exports), labels through the project's merge_map
GH = '/path/to/revision/analyses/original_submission_code'
for fsif in ['3-TF', '3-miR', '3-Comp', '4-TF', '5-TF', '6-TF']:
    Es = set()
    for l in open(f'{GH}/SIF_files/{fsif}.sif'):
        q = l.split()
        if len(q) >= 2: Es.add((merge_map.get(q[0], q[0]), merge_map.get(q[-1], q[-1])))
    tm_ = {(a, b) for a, b in Es if nodes.get(a) == 'TF' and nodes.get(b) == 'miRNA'}; mt_ = {(a, b) for a, b in Es if nodes.get(a) == 'miRNA' and nodes.get(b) == 'TF'}
    P(f'Q5(c) deposit SIF_files/{fsif}.sif: arcs {len(Es)}; TF->miRNA {len(tm_)}; miRNA->TF {len(mt_)}; TF->miRNA reciprocated within this file '
      f'{sum(1 for a, b in tm_ if (b, a) in mt_)}; SIF edge labels {sorted({l.split(chr(9))[1] for l in open(f"{GH}/SIF_files/{fsif}.sif") if chr(9) in l})}')
# the builder: code lines that touch the SIF files in the deposit (quoted)
for fn in ['Scripts/03.1_Signed_FFL_Network_Analysis.R', 'Scripts/Network_Analysis_Legacy.R', 'Scripts/03_Network_Analysis_Hub_Enrichment.R',
           'Scripts/01_TCGA_PanCancer_DEX_GSEA_Survival.R', 'Scripts/02_miRNA_Target_GSEA_Analysis.R', 'Scripts/01.1_BRCA_miRNA_DESeq2_Preprocessing.R']:
    L_ = open(f'{GH}/{fn}', errors='replace').read().split('\n')
    hits = [(i + 1, x.strip()) for i, x in enumerate(L_) if re.search(r'(?i)\.sif|sif_file|writeLines|write\.table|write\.csv|write_tsv', x)]
    P(f'Q5(c) deposit {fn}: lines mentioning SIF files or writing tables: ' + ' || '.join(f'{n}: {x[:110]}' for n, x in hits[:8]))
open(f'{H}/q5_filter_vs_network.txt', 'w').write('\n'.join(out) + '\n')
